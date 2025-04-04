#!/usr/bin/env python3
"""
Enhanced Excel service implementation for spreadsheet creation, reading and manipulation.
"""
import os
import json
import logging
import io
from typing import Dict, List, Any, Optional, Union, Tuple
import pandas as pd

from app.tools.base.service import ToolServiceBase

class ExcelService(ToolServiceBase):
    """Service to handle Excel file creation, reading and manipulation"""

    def __init__(self):
        """Initialize the Excel service"""
        super().__init__()
        self.xlsxwriter = None
        self.pandas = None
        
        # Dictionary to store active workbooks and worksheets
        self.workbooks = {}
        
        # Storage for DataFrames in memory
        self.dataframes = {}
        
        # Flags for optional dependencies
        self.pandas_available = False
        self.openpyxl_available = False
        self.xlrd_available = False
    
    def initialize(self) -> bool:
        """
        Initialize the Excel service with required libraries.
        
        Returns:
            bool: True if initialization was successful, False otherwise
        """
        try:
            # Import required libraries
            import xlsxwriter
            self.xlsxwriter = xlsxwriter
            
            import pandas as pd
            self.pandas = pd
            self.pandas_available = True
            
            # Check for optional dependencies
            try:
                import openpyxl
                self.openpyxl_available = True
            except ImportError:
                self.logger.warning(
                    "openpyxl library not installed. Some Excel reading features may be limited. Install with 'pip install openpyxl'")
                self.openpyxl_available = False
            
            try:
                import xlrd
                self.xlrd_available = True
            except ImportError:
                self.logger.warning(
                    "xlrd library not installed. Reading older Excel formats may be limited. Install with 'pip install xlrd'")
                self.xlrd_available = False
            
            self.logger.info("Excel service initialized successfully")
            self.initialized = True
            return True
            
        except ImportError as e:
            if 'xlsxwriter' in str(e):
                self.logger.error("xlsxwriter library not installed. Please install with 'pip install XlsxWriter'")
            elif 'pandas' in str(e):
                self.logger.error("pandas library not installed. Please install with 'pip install pandas'")
            else:
                self.logger.error(f"Error importing required libraries: {e}")
            return False
        except Exception as e:
            self.logger.error(f"Failed to initialize Excel service: {e}")
            return False
    
    def _check_pandas_available(self):
        """
        Check if pandas is available for DataFrame operations.
        
        Raises:
            ValueError: If pandas is not available
        """
        if not self.pandas_available:
            raise ValueError("pandas library not installed. Please install with 'pip install pandas'")
        return True
    
    # DataFrame storage methods
    
    def store_dataframe(self, dataframe_id: str, df: pd.DataFrame) -> str:
        """
        Store DataFrame in memory for future operations.
        
        Args:
            dataframe_id: ID to use for storing the DataFrame
            df: DataFrame to store
            
        Returns:
            The dataframe_id for reference
        """
        self.dataframes[dataframe_id] = df
        return dataframe_id
    
    def get_dataframe(self, dataframe_id: str) -> Optional[pd.DataFrame]:
        """
        Retrieve DataFrame from memory.
        
        Args:
            dataframe_id: ID of the DataFrame to retrieve
            
        Returns:
            The DataFrame or None if not found
        """
        return self.dataframes.get(dataframe_id)
    
    def list_dataframes(self) -> List[str]:
        """
        List all available DataFrames in memory.
        
        Returns:
            List of DataFrame IDs
        """
        return list(self.dataframes.keys())
    
    def clear_dataframe(self, dataframe_id: str) -> bool:
        """
        Remove DataFrame from memory.
        
        Args:
            dataframe_id: ID of the DataFrame to remove
            
        Returns:
            True if successful, False if DataFrame not found
        """
        if dataframe_id in self.dataframes:
            del self.dataframes[dataframe_id]
            return True
        return False
    
    # Workbook creation methods
    
    def create_workbook(self, filename: str) -> Dict[str, Any]:
        """
        Create a new Excel workbook.
        
        Args:
            filename: Path to save the Excel file
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Create the workbook
            workbook = self.xlsxwriter.Workbook(filename)

            # Store in our dictionary
            self.workbooks[filename] = {
                "workbook": workbook,
                "worksheets": {},
                "formats": {},
                "charts": {}
            }

            return {"filename": filename, "status": "created"}
        except Exception as e:
            self.logger.error(f"Error creating workbook: {e}")
            return {"error": f"Error creating workbook: {str(e)}"}
    
    def add_worksheet(self, filename: str, name: Optional[str] = None) -> Dict[str, Any]:
        """
        Add a worksheet to the workbook.
        
        Args:
            filename: Path to the Excel file
            name: Optional name for the worksheet
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            workbook = self.workbooks[filename]["workbook"]

            # Add a worksheet
            worksheet = workbook.add_worksheet(name)

            # Store the worksheet
            worksheet_name = name if name else f"Sheet{len(self.workbooks[filename]['worksheets'])+1}"
            self.workbooks[filename]["worksheets"][worksheet_name] = worksheet

            return {"filename": filename, "worksheet": worksheet_name, "status": "added"}
        except Exception as e:
            self.logger.error(f"Error adding worksheet: {e}")
            return {"error": f"Error adding worksheet: {str(e)}"}
    
    def write_data(self, filename: str, worksheet_name: str, row: int, col: int, 
                  data: Any, format_name: Optional[str] = None) -> Dict[str, Any]:
        """
        Write data to a cell.
        
        Args:
            filename: Path to the Excel file
            worksheet_name: Name of the worksheet
            row: Row number (0-based)
            col: Column number (0-based)
            data: Data to write
            format_name: Optional name of a predefined format
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            # Check if worksheet exists
            if worksheet_name not in self.workbooks[filename]["worksheets"]:
                return {"error": f"Worksheet {worksheet_name} not found in {filename}"}

            worksheet = self.workbooks[filename]["worksheets"][worksheet_name]

            # Get format if specified
            format_obj = None
            if format_name and format_name in self.workbooks[filename]["formats"]:
                format_obj = self.workbooks[filename]["formats"][format_name]

            # Write the data
            worksheet.write(row, col, data, format_obj)

            return {"filename": filename, "worksheet": worksheet_name, "row": row, "col": col, "data": data, "status": "written"}
        except Exception as e:
            self.logger.error(f"Error writing data: {e}")
            return {"error": f"Error writing data: {str(e)}"}
    
    def write_matrix(self, filename: str, worksheet_name: str, start_row: int, start_col: int,
                    data: List[List[Any]], formats: Optional[List[List[str]]] = None) -> Dict[str, Any]:
        """
        Write a matrix of data to a worksheet.
        
        Args:
            filename: Path to the Excel file
            worksheet_name: Name of the worksheet
            start_row: Starting row number (0-based)
            start_col: Starting column number (0-based)
            data: 2D list of data to write
            formats: Optional 2D list of format names corresponding to data
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            # Check if worksheet exists
            if worksheet_name not in self.workbooks[filename]["worksheets"]:
                return {"error": f"Worksheet {worksheet_name} not found in {filename}"}

            worksheet = self.workbooks[filename]["worksheets"][worksheet_name]

            # Process formats if provided
            format_matrix = None
            if formats:
                format_matrix = []
                for row in formats:
                    format_row = []
                    for format_name in row:
                        if format_name and format_name in self.workbooks[filename]["formats"]:
                            format_row.append(
                                self.workbooks[filename]["formats"][format_name])
                        else:
                            format_row.append(None)
                    format_matrix.append(format_row)

            # Write the data matrix
            for i, row_data in enumerate(data):
                for j, cell_data in enumerate(row_data):
                    format_obj = None
                    if format_matrix and i < len(format_matrix) and j < len(format_matrix[i]):
                        format_obj = format_matrix[i][j]

                    worksheet.write(start_row + i, start_col + j, cell_data, format_obj)

            return {"filename": filename, "worksheet": worksheet_name, "rows": len(data), "cols": len(data[0]) if data else 0, "status": "written"}
        except Exception as e:
            self.logger.error(f"Error writing matrix: {e}")
            return {"error": f"Error writing matrix: {str(e)}"}
    
    def add_format(self, filename: str, format_name: str, format_props: Dict[str, Any]) -> Dict[str, Any]:
        """
        Create a cell format.
        
        Args:
            filename: Path to the Excel file
            format_name: Name to identify the format
            format_props: Dictionary of format properties
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            workbook = self.workbooks[filename]["workbook"]

            # Create the format
            format_obj = workbook.add_format(format_props)

            # Store the format
            self.workbooks[filename]["formats"][format_name] = format_obj

            return {"filename": filename, "format": format_name, "properties": format_props, "status": "added"}
        except Exception as e:
            self.logger.error(f"Error adding format: {e}")
            return {"error": f"Error adding format: {str(e)}"}
    
    def add_chart(self, filename: str, worksheet_name: str, chart_type: str, 
                 data_range: List[Dict[str, Any]], position: Dict[str, int], 
                 options: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """
        Add a chart to a worksheet.
        
        Args:
            filename: Path to the Excel file
            worksheet_name: Name of the worksheet
            chart_type: Type of chart
            data_range: List of data series specifications
            position: Dictionary with row and col keys specifying chart position
            options: Optional additional chart options
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            # Check if worksheet exists
            if worksheet_name not in self.workbooks[filename]["worksheets"]:
                return {"error": f"Worksheet {worksheet_name} not found in {filename}"}

            workbook = self.workbooks[filename]["workbook"]
            worksheet = self.workbooks[filename]["worksheets"][worksheet_name]

            # Create the chart
            chart = workbook.add_chart({'type': chart_type})

            # Add the data series
            for series in data_range:
                chart.add_series(series)

            # Set chart title and other options if provided
            if options:
                if 'title' in options:
                    chart.set_title({'name': options['title']})
                if 'x_axis' in options:
                    chart.set_x_axis(options['x_axis'])
                if 'y_axis' in options:
                    chart.set_y_axis(options['y_axis'])
                if 'style' in options:
                    chart.set_style(options['style'])

            # Insert the chart into the worksheet
            worksheet.insert_chart(position['row'], position['col'], chart)

            # Store the chart
            chart_name = f"Chart{len(self.workbooks[filename]['charts'])+1}"
            self.workbooks[filename]["charts"][chart_name] = chart

            return {"filename": filename, "worksheet": worksheet_name, "chart": chart_name, "status": "added"}
        except Exception as e:
            self.logger.error(f"Error adding chart: {e}")
            return {"error": f"Error adding chart: {str(e)}"}
    
    def add_image(self, filename: str, worksheet_name: str, image_path: str,
                 position: Dict[str, int], options: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """
        Add an image to a worksheet.
        
        Args:
            filename: Path to the Excel file
            worksheet_name: Name of the worksheet
            image_path: Path to the image file
            position: Dictionary with row and col keys specifying image position
            options: Optional additional image options
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            # Check if worksheet exists
            if worksheet_name not in self.workbooks[filename]["worksheets"]:
                return {"error": f"Worksheet {worksheet_name} not found in {filename}"}

            # Check if image file exists
            if not os.path.exists(image_path):
                return {"error": f"Image file {image_path} not found"}

            worksheet = self.workbooks[filename]["worksheets"][worksheet_name]

            # Insert the image
            worksheet.insert_image(position['row'], position['col'], image_path, options)

            return {"filename": filename, "worksheet": worksheet_name, "image": image_path, "status": "added"}
        except Exception as e:
            self.logger.error(f"Error adding image: {e}")
            return {"error": f"Error adding image: {str(e)}"}
    
    def add_formula(self, filename: str, worksheet_name: str, row: int, col: int,
                   formula: str, format_name: Optional[str] = None) -> Dict[str, Any]:
        """
        Add a formula to a cell.
        
        Args:
            filename: Path to the Excel file
            worksheet_name: Name of the worksheet
            row: Row number (0-based)
            col: Column number (0-based)
            formula: Excel formula
            format_name: Optional name of a predefined format
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            # Check if worksheet exists
            if worksheet_name not in self.workbooks[filename]["worksheets"]:
                return {"error": f"Worksheet {worksheet_name} not found in {filename}"}

            worksheet = self.workbooks[filename]["worksheets"][worksheet_name]

            # Get format if specified
            format_obj = None
            if format_name and format_name in self.workbooks[filename]["formats"]:
                format_obj = self.workbooks[filename]["formats"][format_name]

            # Write the formula
            worksheet.write_formula(row, col, formula, format_obj)

            return {"filename": filename, "worksheet": worksheet_name, "row": row, "col": col, "formula": formula, "status": "added"}
        except Exception as e:
            self.logger.error(f"Error adding formula: {e}")
            return {"error": f"Error adding formula: {str(e)}"}
    
    def add_table(self, filename: str, worksheet_name: str, start_row: int, start_col: int,
                 end_row: int, end_col: int, options: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """
        Add a table to a worksheet.
        
        Args:
            filename: Path to the Excel file
            worksheet_name: Name of the worksheet
            start_row: Starting row number (0-based)
            start_col: Starting column number (0-based)
            end_row: Ending row number (0-based)
            end_col: Ending column number (0-based)
            options: Optional table options
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            # Check if worksheet exists
            if worksheet_name not in self.workbooks[filename]["worksheets"]:
                return {"error": f"Worksheet {worksheet_name} not found in {filename}"}

            worksheet = self.workbooks[filename]["worksheets"][worksheet_name]

            # Add the table
            table_options = options or {}
            worksheet.add_table(start_row, start_col, end_row, end_col, table_options)

            return {"filename": filename, "worksheet": worksheet_name, "table_range": f"{start_row}:{start_col}:{end_row}:{end_col}", "status": "added"}
        except Exception as e:
            self.logger.error(f"Error adding table: {e}")
            return {"error": f"Error adding table: {str(e)}"}
    
    def close_workbook(self, filename: str) -> Dict[str, Any]:
        """
        Close and save the workbook.
        
        Args:
            filename: Path to the Excel file
            
        Returns:
            Dictionary with operation result
        """
        self._is_initialized()
        
        try:
            # Check if workbook exists
            if filename not in self.workbooks:
                return {"error": f"Workbook {filename} not found"}

            workbook = self.workbooks[filename]["workbook"]

            # Close the workbook (which saves it)
            workbook.close()

            # Remove from our dictionary
            del self.workbooks[filename]

            return {"filename": filename, "status": "closed"}
        except Exception as e:
            self.logger.error(f"Error closing workbook: {e}")
            return {"error": f"Error closing workbook: {str(e)}"}
    
    # DataFrame handling methods
    
    def read_excel(self, filename: str, sheet_name: Union[str, int] = 0, **kwargs) -> Union[pd.DataFrame, Dict[str, Any]]:
        """
        Read Excel file into DataFrame.
        
        Args:
            filename: Path to the Excel file
            sheet_name: Sheet name or index
            **kwargs: Additional parameters for pandas.read_excel
            
        Returns:
            DataFrame or Dictionary with error information
        """
        self._check_pandas_available()
        
        try:
            # Check if file exists
            if not os.path.exists(filename):
                return {"error": f"File {filename} not found"}

            # Read Excel file
            df = self.pandas.read_excel(filename, sheet_name=sheet_name, **kwargs)

            # If it returned a dict of DataFrames (multiple sheets)
            if isinstance(df, dict):
                return {"sheets": list(df.keys()), "dataframes": df}

            return df
        except Exception as e:
            self.logger.error(f"Error reading Excel file: {e}")
            return {"error": f"Error reading Excel file: {str(e)}"}
    
    def read_csv(self, filename: str, **kwargs) -> Union[pd.DataFrame, Dict[str, Any]]:
        """
        Read CSV file into DataFrame.
        
        Args:
            filename: Path to the CSV file
            **kwargs: Additional parameters for pandas.read_csv
            
        Returns:
            DataFrame or Dictionary with error information
        """
        self._check_pandas_available()
        
        try:
            # Check if file exists
            if not os.path.exists(filename):
                return {"error": f"File {filename} not found"}

            # Read CSV file
            df = self.pandas.read_csv(filename, **kwargs)
            return df
        except Exception as e:
            self.logger.error(f"Error reading CSV file: {e}")
            return {"error": f"Error reading CSV file: {str(e)}"}
    
    def get_excel_sheet_names(self, filename: str) -> Union[List[str], Dict[str, Any]]:
        """
        Get sheet names from Excel file.
        
        Args:
            filename: Path to the Excel file
            
        Returns:
            List of sheet names or Dictionary with error information
        """
        self._check_pandas_available()
        
        try:
            # Check if file exists
            if not os.path.exists(filename):
                return {"error": f"File {filename} not found"}

            # Get sheet names
            excel_file = self.pandas.ExcelFile(filename)
            sheet_names = excel_file.sheet_names

            return sheet_names
        except Exception as e:
            self.logger.error(f"Error getting sheet names: {e}")
            return {"error": f"Error getting sheet names: {str(e)}"}
    
    def dataframe_info(self, df: pd.DataFrame) -> Dict[str, Any]:
        """
        Get information about DataFrame.
        
        Args:
            df: DataFrame to analyze
            
        Returns:
            Dictionary with DataFrame information
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Create a string buffer to capture DataFrame.info() output
            buffer = io.StringIO()
            df.info(buf=buffer)
            info_output = buffer.getvalue()

            # Basic info about the DataFrame
            result = {
                "shape": df.shape,
                "columns": df.columns.tolist(),
                # Convert to strings for JSON serialization
                "dtypes": {str(k): str(v) for k, v in df.dtypes.to_dict().items()},
                "info": info_output,
                "memory_usage": df.memory_usage(deep=True).sum(),
                "has_nulls": df.isnull().any().any()
            }

            # Add sample data (first 5 rows)
            try:
                result["head"] = json.loads(df.head().to_json(orient="records"))
            except:
                # Fallback if to_json fails for some reason
                result["head"] = str(df.head())

            return result
        except Exception as e:
            self.logger.error(f"Error getting DataFrame info: {e}")
            return {"error": f"Error getting DataFrame info: {str(e)}"}
    
    def dataframe_to_excel(self, df: pd.DataFrame, filename: str, 
                         sheet_name: str = "Sheet1", index: bool = True, **kwargs) -> Dict[str, Any]:
        """
        Export DataFrame to Excel file.
        
        Args:
            df: DataFrame to export
            filename: Path to save the Excel file
            sheet_name: Name of the sheet
            index: Whether to include the DataFrame index
            **kwargs: Additional parameters for DataFrame.to_excel
            
        Returns:
            Dictionary with operation result
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Export to Excel
            df.to_excel(filename, sheet_name=sheet_name, index=index, **kwargs)

            return {"filename": filename, "sheet_name": sheet_name, "rows": len(df), "columns": len(df.columns), "status": "exported"}
        except Exception as e:
            self.logger.error(f"Error exporting to Excel: {e}")
            return {"error": f"Error exporting to Excel: {str(e)}"}
    
    def dataframe_to_csv(self, df: pd.DataFrame, filename: str, 
                      index: bool = True, **kwargs) -> Dict[str, Any]:
        """
        Export DataFrame to CSV file.
        
        Args:
            df: DataFrame to export
            filename: Path to save the CSV file
            index: Whether to include the DataFrame index
            **kwargs: Additional parameters for DataFrame.to_csv
            
        Returns:
            Dictionary with operation result
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Export to CSV
            df.to_csv(filename, index=index, **kwargs)

            return {"filename": filename, "rows": len(df), "columns": len(df.columns), "status": "exported"}
        except Exception as e:
            self.logger.error(f"Error exporting to CSV: {e}")
            return {"error": f"Error exporting to CSV: {str(e)}"}
    
    def filter_dataframe(self, df: pd.DataFrame, query: Optional[str] = None, 
                      column: Optional[str] = None, value: Any = None, 
                      operator: str = "==") -> Union[pd.DataFrame, Dict[str, Any]]:
        """
        Filter DataFrame by query or column condition.
        
        Args:
            df: DataFrame to filter
            query: Query string for filtering
            column: Column name to filter by (alternative to query)
            value: Value to compare with (used with column and operator)
            operator: Comparison operator (==, !=, >, >=, <, <=, in, contains)
            
        Returns:
            Filtered DataFrame or Dictionary with error information
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Filter by query string
            if query:
                try:
                    filtered_df = df.query(query)
                    return filtered_df
                except Exception as e:
                    self.logger.error(f"Error in query: {e}")
                    return {"error": f"Error in query: {str(e)}"}

            # Filter by column and value
            elif column and value is not None:
                if column not in df.columns:
                    return {"error": f"Column '{column}' not found"}

                if operator == "==":
                    filtered_df = df[df[column] == value]
                elif operator == "!=":
                    filtered_df = df[df[column] != value]
                elif operator == ">":
                    filtered_df = df[df[column] > value]
                elif operator == ">=":
                    filtered_df = df[df[column] >= value]
                elif operator == "<":
                    filtered_df = df[df[column] < value]
                elif operator == "<=":
                    filtered_df = df[df[column] <= value]
                elif operator == "in":
                    if not isinstance(value, list):
                        return {"error": "Value must be a list when using 'in' operator"}
                    filtered_df = df[df[column].isin(value)]
                elif operator == "contains":
                    filtered_df = df[df[column].astype(str).str.contains(str(value))]
                else:
                    return {"error": f"Unknown operator: {operator}"}

                return filtered_df
            else:
                return {"error": "Either query or column+value must be provided"}

        except Exception as e:
            self.logger.error(f"Error filtering DataFrame: {e}")
            return {"error": f"Error filtering DataFrame: {str(e)}"}
    
    def sort_dataframe(self, df: pd.DataFrame, by: Union[str, List[str]], 
                     ascending: Union[bool, List[bool]] = True) -> Union[pd.DataFrame, Dict[str, Any]]:
        """
        Sort DataFrame by columns.
        
        Args:
            df: DataFrame to sort
            by: Column name(s) to sort by
            ascending: Whether to sort in ascending order
            
        Returns:
            Sorted DataFrame or Dictionary with error information
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Ensure 'by' is a list
            if isinstance(by, str):
                by = [by]

            # Check if all columns exist
            for col in by:
                if col not in df.columns:
                    return {"error": f"Column '{col}' not found"}

            # Convert ascending to list if necessary
            if not isinstance(ascending, list):
                ascending = [ascending] * len(by)

            # Sort the DataFrame
            sorted_df = df.sort_values(by=by, ascending=ascending)

            return sorted_df
        except Exception as e:
            self.logger.error(f"Error sorting DataFrame: {e}")
            return {"error": f"Error sorting DataFrame: {str(e)}"}
    
    def group_dataframe(self, df: pd.DataFrame, by: Union[str, List[str]], 
                      agg_func: Any = None) -> Union[pd.DataFrame, Dict[str, Any]]:
        """
        Group DataFrame and apply aggregation.
        
        Args:
            df: DataFrame to group
            by: Column name(s) to group by
            agg_func: Aggregation function(s) to apply
            
        Returns:
            Grouped DataFrame or Dictionary with error information
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Ensure 'by' is a list
            if isinstance(by, str):
                by = [by]

            # Check if all columns exist
            for col in by:
                if col not in df.columns:
                    return {"error": f"Column '{col}' not found"}

            # Group by columns
            grouped = df.groupby(by)

            # Apply aggregation function
            if agg_func:
                if isinstance(agg_func, dict):
                    # Complex aggregation with different functions per column
                    result_df = grouped.agg(agg_func)
                else:
                    # Simple aggregation with same function for all columns
                    result_df = grouped.agg(agg_func)
            else:
                # Default aggregation (count)
                result_df = grouped.count()

            return result_df.reset_index()
        except Exception as e:
            self.logger.error(f"Error grouping DataFrame: {e}")
            return {"error": f"Error grouping DataFrame: {str(e)}"}
    
    def describe_dataframe(self, df: pd.DataFrame, 
                        include: Any = None, 
                        exclude: Any = None, 
                        percentiles: Optional[List[float]] = None) -> Dict[str, Any]:
        """
        Get statistical description of DataFrame.
        
        Args:
            df: DataFrame to describe
            include: Types of columns to include
            exclude: Types of columns to exclude
            percentiles: List of percentiles to include in output
            
        Returns:
            Dictionary with statistical description
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Get description
            desc_df = df.describe(include=include, exclude=exclude, percentiles=percentiles)

            # Convert to dict for JSON serialization
            result = {}
            for col in desc_df.columns:
                result[str(col)] = {str(idx): float(val) if not pd.isna(val) else None
                                    for idx, val in desc_df[col].items()}

            return result
        except Exception as e:
            self.logger.error(f"Error describing DataFrame: {e}")
            return {"error": f"Error describing DataFrame: {str(e)}"}
    
    def get_column_values(self, df: pd.DataFrame, column: str, 
                        unique: bool = False, count: bool = False) -> Dict[str, Any]:
        """
        Get values from a specific column in a DataFrame.
        
        Args:
            df: DataFrame to analyze
            column: Name of the column to get values from
            unique: Whether to return only unique values
            count: Whether to count occurrences of each value
            
        Returns:
            Dictionary with column values or value counts
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Check if column exists
            if column not in df.columns:
                return {"error": f"Column '{column}' not found"}

            if count:
                # Count occurrences of each value
                value_counts = df[column].value_counts().to_dict()

                return {
                    "column": column,
                    "value_counts": value_counts,
                    "total_values": len(df[column]),
                    "unique_values": len(value_counts),
                    "status": "success"
                }

            elif unique:
                # Get unique values
                unique_values = df[column].unique().tolist()

                # Convert non-serializable values to strings if necessary
                unique_values = [str(v) if not isinstance(v, (int, float, bool, str, type(None))) else v for v in unique_values]

                return {
                    "column": column,
                    "unique_values": unique_values,
                    "count": len(unique_values),
                    "status": "success"
                }

            else:
                # Get all values
                values = df[column].tolist()

                # Convert non-serializable values to strings if necessary
                values = [str(v) if not isinstance(v, (int, float, bool, str, type(None))) else v for v in values]

                return {
                    "column": column,
                    "values": values,
                    "count": len(values),
                    "status": "success"
                }
        except Exception as e:
            self.logger.error(f"Error getting column values: {e}")
            return {"error": f"Error getting column values: {str(e)}"}
    
    def get_correlation(self, df: pd.DataFrame, method: str = "pearson") -> Dict[str, Any]:
        """
        Get correlation matrix for a DataFrame.
        
        Args:
            df: DataFrame to analyze
            method: Correlation method ('pearson', 'kendall', or 'spearman')
            
        Returns:
            Dictionary with correlation matrix
        """
        self._check_pandas_available()
        
        try:
            # Check if valid DataFrame
            if not isinstance(df, self.pandas.DataFrame):
                return {"error": "Invalid DataFrame"}

            # Calculate correlation matrix
            corr = df.corr(method=method).round(4).to_dict()

            return {
                "correlation_matrix": corr,
                "method": method,
                "status": "success"
            }
        except Exception as e:
            self.logger.error(f"Error calculating correlation: {e}")
            return {"error": f"Error calculating correlation: {str(e)}"}
